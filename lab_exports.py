"""Markdown-ready PDF and Word exports for synthesized lab protocols.

The renderer intentionally supports the small Markdown subset produced by the
synthesis engine: headings, paragraphs, bold/italic/code spans, bullet and
numbered lists, block quotes, horizontal rules and pipe tables.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from html import escape


@dataclass(frozen=True)
class Block:
    kind: str
    text: str = ""
    level: int = 0
    rows: tuple[tuple[str, ...], ...] = field(default_factory=tuple)


_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
_QUOTE = re.compile(r"^>\s?(.*)$")
_RULE = re.compile(r"^\s*(?:---+|\*\*\*+|___+)\s*$")
_INLINE = re.compile(r"(\*\*.+?\*\*|__.+?__|\*[^*\n]+?\*|_[^_\n]+?_|`[^`\n]+?`|\[[^\]]+\]\([^)]+\))")


def _is_table_separator(line: str) -> bool:
    stripped = line.strip().strip("|")
    if not stripped:
        return False
    cells = [cell.strip() for cell in stripped.split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", cell) for cell in cells)


def _split_row(line: str) -> tuple[str, ...]:
    stripped = line.strip()
    stripped = stripped.removeprefix("|")
    stripped = stripped.removesuffix("|")
    return tuple(cell.strip() for cell in stripped.split("|"))


def parse_markdown(markdown: str) -> list[Block]:
    """Parse the protocol Markdown into renderer-neutral blocks."""
    lines = markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[Block] = []
    paragraph: list[str] = []
    index = 0

    def flush() -> None:
        if paragraph:
            blocks.append(Block("paragraph", " ".join(part.strip() for part in paragraph)))
            paragraph.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            flush()
            index += 1
            continue
        if stripped.startswith("```"):
            flush()
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            blocks.append(Block("code", "\n".join(code)))
            index += 1
            continue
        heading = _HEADING.match(stripped)
        if heading:
            flush()
            blocks.append(Block("heading", heading.group(2).strip(), len(heading.group(1))))
            index += 1
            continue
        if _RULE.match(stripped):
            flush()
            blocks.append(Block("rule"))
            index += 1
            continue
        if "|" in stripped and index + 1 < len(lines) and _is_table_separator(lines[index + 1]):
            flush()
            rows: list[tuple[str, ...]] = [_split_row(stripped)]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(_split_row(lines[index]))
                index += 1
            width = max(len(row) for row in rows)
            rows = [row + ("",) * (width - len(row)) for row in rows]
            blocks.append(Block("table", rows=tuple(rows)))
            continue
        bullet = _BULLET.match(line)
        if bullet:
            flush()
            blocks.append(Block("bullet", bullet.group(2).strip(),
                                len(bullet.group(1).replace("\t", "  ")) // 2 + 1))
            index += 1
            continue
        numbered = _NUMBERED.match(line)
        if numbered:
            flush()
            blocks.append(Block("numbered", numbered.group(2).strip(),
                                len(numbered.group(1).replace("\t", "  ")) // 2 + 1))
            index += 1
            continue
        quote = _QUOTE.match(stripped)
        if quote:
            flush()
            blocks.append(Block("quote", quote.group(1).strip()))
            index += 1
            continue
        paragraph.append(line)
        index += 1
    flush()
    return blocks


def _inline_runs(text: str) -> list[tuple[str, bool, bool, bool]]:
    """Split text into ``(text, bold, italic, code)`` runs."""
    runs: list[tuple[str, bool, bool, bool]] = []
    position = 0
    for match in _INLINE.finditer(text):
        if match.start() > position:
            runs.append((text[position:match.start()], False, False, False))
        token = match.group(0)
        if token.startswith(("**", "__")):
            runs.append((token[2:-2], True, False, False))
        elif token.startswith("`"):
            runs.append((token[1:-1], False, False, True))
        elif token.startswith("["):
            label, url = token[1:-1].split("](", 1)
            runs.append((f"{label} ({url})", False, False, False))
        else:
            runs.append((token[1:-1], False, True, False))
        position = match.end()
    if position < len(text):
        runs.append((text[position:], False, False, False))
    return runs or [(text, False, False, False)]


def _plain(text: str) -> str:
    return "".join(part for part, _, _, _ in _inline_runs(text))


def _pdf_inline(text: str) -> str:
    pieces = []
    for part, bold, italic, code in _inline_runs(text):
        escaped = escape(part)
        if code:
            escaped = f'<font face="Courier">{escaped}</font>'
        if bold:
            escaped = f"<b>{escaped}</b>"
        if italic:
            escaped = f"<i>{escaped}</i>"
        pieces.append(escaped)
    return "".join(pieces)


def docx_bytes(markdown: str) -> bytes:
    """Render Markdown as a Word-compatible .docx document."""
    from docx import Document
    from docx.shared import Cm, Pt

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    def add_runs(paragraph, text: str) -> None:
        for part, bold, italic, code in _inline_runs(text):
            run = paragraph.add_run(part)
            run.bold = bold
            run.italic = italic
            if code:
                run.font.name = "Consolas"

    for block in parse_markdown(markdown):
        if block.kind == "heading":
            document.add_heading(_plain(block.text), level=min(max(block.level, 1), 4))
        elif block.kind == "paragraph":
            add_runs(document.add_paragraph(), block.text)
        elif block.kind == "bullet":
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Cm(0.6 * max(block.level - 1, 0))
            add_runs(paragraph, block.text)
        elif block.kind == "numbered":
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.left_indent = Cm(0.6 * max(block.level - 1, 0))
            add_runs(paragraph, block.text)
        elif block.kind == "quote":
            paragraph = document.add_paragraph(style="Intense Quote")
            add_runs(paragraph, block.text)
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif block.kind == "rule":
            document.add_paragraph("─" * 40)
        elif block.kind == "table" and block.rows:
            table = document.add_table(rows=len(block.rows), cols=len(block.rows[0]))
            table.style = "Table Grid"
            for row_index, row in enumerate(block.rows):
                for column_index, cell in enumerate(row):
                    paragraph = table.cell(row_index, column_index).paragraphs[0]
                    add_runs(paragraph, cell)
                    if row_index == 0:
                        for run in paragraph.runs:
                            run.bold = True
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_bytes(markdown: str) -> bytes:
    """Render Markdown as an A4 PDF with tables, lists and page numbers."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        HRFlowable,
        KeepTogether,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    output = io.BytesIO()
    page_width = A4[0]
    margin = 18 * mm
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "ProtocolBody", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=9.5, leading=13, alignment=TA_JUSTIFY, spaceAfter=4,
    )
    small = ParagraphStyle(
        "ProtocolCell", parent=body, fontSize=8, leading=10, alignment=0,
        spaceAfter=0,
    )
    bullet = ParagraphStyle(
        "ProtocolBullet", parent=body, leftIndent=12, bulletIndent=3, spaceAfter=2,
    )
    code_style = ParagraphStyle(
        "ProtocolCode", parent=body, fontName="Courier", fontSize=8,
        leading=10, backColor=colors.HexColor("#f4f4f4"), leftIndent=6,
    )

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillGray(0.45)
        canvas.drawRightString(page_width - margin, 10 * mm, str(doc.page))
        canvas.restoreState()

    doc = SimpleDocTemplate(
        output, pagesize=A4, leftMargin=margin, rightMargin=margin,
        topMargin=margin, bottomMargin=margin,
    )
    story = []
    for block in parse_markdown(markdown):
        if block.kind == "heading":
            level = min(max(block.level, 1), 4)
            style = ParagraphStyle(
                f"ProtocolH{level}", parent=styles["Heading2"] if level > 1 else styles["Title"],
                fontName="Helvetica-Bold", fontSize={1: 18, 2: 13, 3: 11, 4: 10}[level],
                leading={1: 22, 2: 17, 3: 14, 4: 13}[level],
                spaceBefore={1: 0, 2: 10, 3: 7, 4: 5}[level], spaceAfter=4,
            )
            story.append(Paragraph(_pdf_inline(block.text), style))
        elif block.kind == "paragraph":
            story.append(Paragraph(_pdf_inline(block.text), body))
        elif block.kind in {"bullet", "numbered"}:
            marker = "•" if block.kind == "bullet" else "1."
            story.append(Paragraph(
                _pdf_inline(block.text), bullet,
                bulletText=marker,
            ))
        elif block.kind == "quote":
            story.append(Paragraph(f"<i>{_pdf_inline(block.text)}</i>", ParagraphStyle(
                "ProtocolQuote", parent=body, leftIndent=14, textColor=colors.HexColor("#555555"),
            )))
        elif block.kind == "code":
            story.append(Paragraph(escape(block.text).replace("\n", "<br/>"), code_style))
        elif block.kind == "rule":
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#cccccc")))
            story.append(Spacer(1, 4))
        elif block.kind == "table" and block.rows:
            header = [Paragraph(f"<b>{_pdf_inline(cell)}</b>", small) for cell in block.rows[0]]
            rows = [header] + [
                [Paragraph(_pdf_inline(cell), small) for cell in row]
                for row in block.rows[1:]
            ]
            available = page_width - 2 * margin
            column_width = max(available / max(len(block.rows[0]), 1), 18)
            table = Table(rows, colWidths=[column_width] * len(block.rows[0]), repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef5")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b9c4d0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f9fb")]),
            ]))
            story.append(KeepTogether([table, Spacer(1, 6)]))
    if not story:
        story.append(Paragraph(" ", body))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
